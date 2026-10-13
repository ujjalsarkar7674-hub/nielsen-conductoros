import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BG      = RGBColor(0x0A, 0x0A, 0x1A)
CARD    = RGBColor(0x12, 0x12, 0x2A)
CARD2   = RGBColor(0x18, 0x18, 0x32)
BORDER  = RGBColor(0x24, 0x24, 0x48)
TEXT    = RGBColor(0xF0, 0xF0, 0xF5)
GREEN   = RGBColor(0x00, 0xE6, 0x76)
BLUE    = RGBColor(0x44, 0x8A, 0xFF)
GRAY    = RGBColor(0x9E, 0x9E, 0xB0)
RED_C   = RGBColor(0xFF, 0x52, 0x52)
AMBER   = RGBColor(0xFF, 0xB3, 0x00)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DKGRN   = RGBColor(0x00, 0x38, 0x18)

FONT = 'Arial'

SW = Inches(13.333)
SH = Inches(7.5)
M  = Inches(0.65)
CW = SW - 2 * M

BASE_DIR  = r'C:\Users\UJJAL_SARKAR\Downloads\Neilson\conductor-os'
QR_PATH   = os.path.join(BASE_DIR, 'submission', 'qr-code.png')
SHOTS_DIR = os.path.join(BASE_DIR, 'submission', 'screenshots')
OUT_PATH  = os.path.join(BASE_DIR, 'submission', 'ConductorOS_Presentation.pptx')

SHOT_GRID    = os.path.join(SHOTS_DIR, '01_main_grid.png')
SHOT_CONFIG  = os.path.join(SHOTS_DIR, '02_ticket_config.png')
SHOT_SUCCESS = os.path.join(SHOTS_DIR, '03_ticket_success.png')
SHOT_UNDO    = os.path.join(SHOTS_DIR, '04_undo_active_grid.png')
SHOT_SUN     = os.path.join(SHOTS_DIR, '05_sunlight_mode.png')

def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def rect(slide, x, y, w, h, bg=None, line=None, lw=0.75):
    shp = slide.shapes.add_shape(1, x, y, w, h)
    if bg:
        shp.fill.solid()
        shp.fill.fore_color.rgb = bg
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    return shp

def textbox(slide, text, x, y, w, h, sz=11, bold=False, italic=False,
            col=None, align=PP_ALIGN.LEFT, wrap=True):
    if col is None:
        col = TEXT
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = col
    return box

def multi_text(slide, lines, x, y, w, h, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ln.get('align', align)
        if 'sb' in ln:
            p.space_before = Pt(ln['sb'])
        r = p.add_run()
        r.text = ln.get('t', '')
        r.font.name = FONT
        r.font.size = Pt(ln.get('sz', 11))
        r.font.bold = ln.get('b', False)
        r.font.italic = ln.get('i', False)
        r.font.color.rgb = ln.get('c', TEXT)
    return box

def bar(slide, x, y, w, col=GREEN, h=Pt(2.5)):
    shp = slide.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = col
    shp.line.fill.background()

def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()

def slide_num(slide, n, total=10):
    textbox(slide, f'{n} / {total}', SW - M - Inches(0.8), SH - Inches(0.42),
            Inches(0.8), Inches(0.3), sz=8, col=GRAY, align=PP_ALIGN.RIGHT)

def headline(slide, title, y=Inches(0.4)):
    textbox(slide, title, M, y, CW, Inches(0.55), sz=23, bold=True, col=WHITE)
    bar(slide, M, y + Inches(0.58), Inches(2.2), GREEN, Pt(3))

def add_link_run(p, text, url, sz=10, col=BLUE, bold=False):
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(sz)
    r.font.color.rgb = col
    r.font.bold = bold
    r.font.underline = True
    r.hyperlink.address = url
    return r

def add_image(slide, path, x, y, w=None, h=None):
    if os.path.exists(path):
        if w and h:
            return slide.shapes.add_picture(path, x, y, width=w, height=h)
        elif w:
            return slide.shapes.add_picture(path, x, y, width=w)
        elif h:
            return slide.shapes.add_picture(path, x, y, height=h)
        else:
            return slide.shapes.add_picture(path, x, y)
    else:
        fw = w or Inches(2.5)
        fh = h or Inches(5.0)
        rect(slide, x, y, fw, fh, bg=CARD, line=BORDER)
        textbox(slide, f'[Image Missing]\n{os.path.basename(path)}', x, y + fh/2 - Inches(0.3),
                fw, Inches(0.6), sz=9, col=GRAY, align=PP_ALIGN.CENTER)
        return None

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])
