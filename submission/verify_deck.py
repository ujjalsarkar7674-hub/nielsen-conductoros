import os
from pptx import Presentation

p = r"C:\Users\UJJAL_SARKAR\Downloads\Neilson\conductor-os\submission\ConductorOS_Presentation.pptx"
prs = Presentation(p)
sz = os.path.getsize(p)

print(f"File Size: {sz} bytes ({round(sz/1024, 1)} KB)")
print(f"Total Slides: {len(prs.slides)}")

for i, s in enumerate(prs.slides):
    notes = s.notes_slide.notes_text_frame.text if s.notes_slide else "NO NOTES"
    img_count = sum(1 for shape in s.shapes if shape.shape_type == 13)
    text_count = sum(1 for shape in s.shapes if shape.has_text_frame)
    first_title = ""
    for shape in s.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            first_title = shape.text_frame.text.split("\n")[0][:45]
            break
    print(f"Slide {i+1}: Shapes={len(s.shapes)}, Images={img_count}, NotesLen={len(notes)} chars, Header=\"{first_title}\"")
